import pandas as pd
from PIL import Image
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.chart import PieChart, PieChart3D
from openpyxl.chart import Reference, Series, BarChart
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
loc = 'C:/Users/allen/Downloads/Resources/' \
      'Python/Lib/L1/Excel/data/2'                              # Location of the directory
wb = Workbook()                                                 # Create workbook
ws = wb.active                                                  # Go to active sheet of wb

data = [
    ['Flavour', 'Number Sold'],
    ['Chocolate', 120],
    ['Vanilla', 30],
    ['Strawberry', 20],
    ['Pistachio', 90],
    ['Mint Chocolate Chip', 40],
    ['Cookies and Cream', 60]
]                                                               # Data to be written
for z in data: ws.append(z)                                     # Write data to the sheet

# Pie Chart
chart1 = PieChart()                                             # Create chart
chart1.title = 'Ice_Cream_Chart_Pie'                            # Title of the chart
labels = Reference(ws, min_col=1, min_row=2, max_row=7)         # Labels of the chart
data = Reference(ws, min_col=2, min_row=1, max_row=7)           # Data of the chart
chart1.add_data(data, titles_from_data=True)                    # Add data to the chart
chart1.set_categories(labels)                                   # Set labels of the chart
chart1.height = 7.5                                             # Height of the chart
chart1.width = 10                                               # Width of the chart
chart1.style = 42                                               # Style of the chart

# Bar Chart
chart2 = BarChart()                                             # Create chart
chart2.title = 'Ice_Cream_Chart_Bar'                            # Title of the chart
labels = Reference(ws, min_col=1, min_row=2, max_row=7)         # Labels of the chart
data = Reference(ws, min_col=2, min_row=1, max_row=7)           # Data of the chart
chart2.add_data(data, titles_from_data=True)                    # Add data to the chart
chart2.set_categories(labels)                                   # Set labels of the chart
chart2.height = 7.5                                             # Height of the chart
chart2.width = 15                                               # Width of the chart
chart2.style = 42                                               # Style of the chart

# Table
tab = Table(displayName='Ice_Cream_Table', ref='A1:B7')         # Create table
tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium10',  # Name of the style from excel
    showFirstColumn=False, showLastColumn=False,
    showRowStripes=False, showColumnStripes=True
)                                                               # Style of the table   

# Image
img = Image(f'{loc}/C_Image.png')                               # Create image
img.width = img.width * 0.32                                    # Width of the image
img.height = img.height * 0.32                                  # Height of the image


# Save
ws.add_chart(chart1, 'D1')                                      # Add chart to the sheet
ws.add_chart(chart2, 'K1')                                      # Add chart to the sheet
ws.add_table(tab)                                               # Add table to the sheet
ws.add_image(img, 'A9')                                         # Add image to the sheet
wb.save(f'{loc}/C_Write.xlsx')                                  # Save the workbook